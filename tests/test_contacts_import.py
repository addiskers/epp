"""contacts_import.py — phone normalisation, xlsx/csv parsing, the sample template."""
import io

from openpyxl import Workbook, load_workbook

import contacts_import as ci


def test_normalize_phone_india_centric():
    assert ci.normalize_phone("9876543210") == ("+919876543210", True)
    assert ci.normalize_phone("09876543210") == ("+919876543210", True)
    assert ci.normalize_phone("+91 98765 43210") == ("+919876543210", True)
    assert ci.normalize_phone("9.17619E+11") == (None, False)          # Excel corruption
    assert ci.normalize_phone("12") == ("+12", False)
    assert ci.normalize_phone("") == (None, False)
    assert ci.normalize_phone(None) == (None, False)


def test_csv_with_headers_and_type_column():
    data = (b"Name,Phone,Type,Notes\n"
            b"Rahul,9876543210,Employee,Halol\n"
            b",9876543210,,\n"
            b"Meera,+919000000001,customer,\n"
            b"bad,9.1E+11,,\n")
    rows, rejected, total, unknown = ci.parse_upload("x.csv", data)
    assert total == 4 and rejected == 1 and unknown == []
    assert rows == [("Rahul", "+919876543210", "valid", {"caller_type": "employee", "notes": "Halol"}),
                    ("Meera", "+919000000001", "valid", {"caller_type": "customer", "notes": ""})]


def test_unknown_headers_are_reported_not_dropped_silently():
    rows, _, _, unknown = ci.parse_upload("x.csv", b"Name,Mobile,Shoe size\nA,9000000001,42\n")
    assert rows[0][1] == "+919000000001" and unknown == ["Shoe size"]


def test_headerless_two_column_sheet_and_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.append(["Jay", "9000000002"])
    ws.append(["", "9000000003"])
    buf = io.BytesIO()
    wb.save(buf)
    rows, rejected, total, unknown = ci.parse_upload("c.xlsx", buf.getvalue())
    assert [r[1] for r in rows] == ["+919000000002", "+919000000003"]
    assert rows[0][0] == "Jay" and rows[1][0] == ""
    assert rows[0][3] == {"caller_type": "", "notes": ""}


def test_template_is_a_workbook_with_text_phone_column():
    wb = load_workbook(io.BytesIO(ci.build_template()))
    ws = wb.active
    assert [c.value for c in ws[1]] == ["Name", "Phone", "Type", "Notes"]
    assert ws["B2"].number_format == "@"


def test_a_spreadsheet_number_cell_that_reads_as_a_whole_float_is_accepted():
    """A plain 10-digit number typed into Excel without +91 is a numeric cell; some exports hand
    it over as '9876543210.0'. That is a real number, not scientific-notation corruption."""
    assert ci.normalize_phone("9876543210.0") == ("+919876543210", True)
    assert ci.normalize_phone(9876543210) == ("+919876543210", True)
    assert ci.normalize_phone(9876543210.0) == ("+919876543210", True)
    assert ci.normalize_phone("9.87654321E+09") == (None, False)      # still rejected
